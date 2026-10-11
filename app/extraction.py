from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import pymupdf  # PyMuPDF
from docx import Document as DocxDocument


class ExtractionError(Exception):
    """A problem we can safely explain to the user."""


@dataclass
class PageText:
    page_number: Optional[int]  # 1-based for PDFs; None for DOCX (Word has no fixed pages)
    text: str


def extract_pages(path: str) -> List[PageText]:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    raise ExtractionError(f"Unsupported file type '{ext}'.")


def order_blocks(blocks: list, page_width: float) -> list:
    """Put text blocks in reading order. `blocks` are PyMuPDF tuples (x0, y0, x1, y1, text, ...) already
    sorted top-to-bottom. On a two-column page that order interleaves the columns line by line, so a
    worked example gets mixed with unrelated text from the other column. Here a page counts as two
    columns only when BOTH halves hold several long paragraphs (short cells of a table never do), and
    then each column is read top to bottom, left first. Blocks that span the gutter (headings, footers)
    stay where they are and separate the page into bands."""
    mid, margin = page_width / 2, page_width * 0.02

    def side(b) -> str:
        if b[0] < mid - margin and b[2] > mid + margin:
            return "full"
        return "left" if (b[0] + b[2]) / 2 < mid else "right"

    long_per_side = {"left": 0, "right": 0}
    for b in blocks:
        if side(b) in long_per_side and len(str(b[4]).strip()) > 120:
            long_per_side[side(b)] += 1
    if min(long_per_side.values()) < 2:
        return list(blocks)

    ordered, band = [], []

    def flush() -> None:
        for wanted in ("left", "right"):
            ordered.extend(sorted((b for b in band if side(b) == wanted), key=lambda b: (b[1], b[0])))
        band.clear()

    for b in blocks:
        if side(b) == "full":
            flush()
            ordered.append(b)
        else:
            band.append(b)
    flush()
    return ordered


def _extract_pdf(path: str) -> List[PageText]:
    with pymupdf.open(path) as pdf:
        pages = []
        for i, page in enumerate(pdf):
            # "blocks" keeps paragraphs and headings apart, so we keep that structure
            raw = [b for b in page.get_text("blocks", sort=True) if b[6] == 0]
            blocks = [b[4] for b in order_blocks(raw, page.rect.width)]
            pages.append(PageText(i + 1, "\n\n".join(blocks)))
        return pages


def _extract_docx(path: str) -> List[PageText]:
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [PageText(None, "\n\n".join(parts))]
